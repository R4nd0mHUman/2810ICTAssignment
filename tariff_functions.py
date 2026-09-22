"""Core calculation and data functions for the XPower prototype."""

from __future__ import annotations # Makes type annotations safer/more flexible.

import csv # Reads CSV electricity data.
from datetime import date, datetime, time # Works with calendar dates, date + time timestamps, and 24hr time.
from pathlib import Path # Handles file paths.
from typing import Iterable # Provides type hints for collections of records.


REQUIRED_COLUMNS = {"timestamp", "kWh"}
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


def _number(name, value, *, allow_zero=True):
    """Return a validated non-negative number as a float.
    Type checks are performed first. bool is explicitly excluded because bool is a subclass of int in Python,
                        which would otherwise allow True/False to be treated as 1/0.
    """
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TypeError(f"{name} must be a number.")

    value = float(value)

    if value < 0 or (not allow_zero and value == 0):
        rule = "greater than zero" if not allow_zero else "non-negative"
        raise ValueError(f"{name} must be {rule}.")

    return value


def _parse_timestamp(value):
    """Convert common spreadsheet/CSV timestamp values into a datetime."""
    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, time.min)

    if not isinstance(value, str) or not value.strip():
        raise ValueError("Every usage record must contain a valid timestamp.")

    text = value.strip()

    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    try:
        return datetime.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"Invalid timestamp: {value!r}.") from exc


def _normalise_rows(rows: Iterable[dict]):
    """Validate imported rows and return timestamp/kWh dictionaries."""
    records = []

    for row_number, row in enumerate(rows, start=2):
        try:
            timestamp = _parse_timestamp(row.get("timestamp"))
            consumption = float(row.get("kWh"))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid data in row {row_number}: {exc}") from exc

        if consumption < 0:
            raise ValueError(
                f"Electricity consumption cannot be negative (row {row_number})."
            )

        records.append({
            "timestamp": timestamp,
            "kWh": consumption,
        })

    if not records:
        raise ValueError("The usage file contains no electricity usage records.")

    return records


def load_usage_records(file_path):
    """Load and validate hourly or daily usage records from CSV or Excel."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Usage data must be supplied as a .csv or .xlsx file.")

    if not path.exists():
        raise FileNotFoundError(f"Usage file not found: {path}")

    if extension == ".csv":
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)

            # Check that the CSV has a header row.
            if reader.fieldnames is None:
                raise ValueError("CSV file is empty or has no header row.")

            # Check that the required columns exist.
            if not REQUIRED_COLUMNS.issubset(reader.fieldnames):
                raise ValueError(
                    "Usage file must contain 'timestamp' and 'kWh' columns."
                )

            return _normalise_rows(reader)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel support requires openpyxl.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header = next(rows, None)

    if header is None:
        raise ValueError("Excel file is empty or has no header row.")

    headers = [
        str(value).strip() if value is not None else ""
        for value in header
    ]

    if not REQUIRED_COLUMNS.issubset(headers):
        raise ValueError(
            "Usage file must contain 'timestamp' and 'kWh' columns."
        )

    mapped = (dict(zip(headers, values)) for values in rows)

    return _normalise_rows(mapped)


def load_total_consumption(file_path):
    """
    Read household electricity usage from a CSV or Excel file and return the total electricity consumption in kWh.

    The usage file must contain:
        - timestamp
        - kWh

    Parameters:
        - file_path: Path to the electricity usage file.

    Returns:
        - float: Total electricity consumption in kWh.

    Raises:
        - ValueError: If the required columns are missing.
        - ValueError: If a kWh value is not a valid number.
        - ValueError: If a kWh value is negative.
        - ValueError: If the file contains no usage records.
        - FileNotFoundError: If the usage file does not exist.
    """
    return round(
        sum(record["kWh"] for record in load_usage_records(file_path)),
        2,
    )


def filter_usage_records(records, start_date=None, end_date=None):
    """Return records within an inclusive user-selected date range."""
    if not records:
        raise ValueError("At least one usage record is required.")

    start = _parse_timestamp(start_date).date() if start_date else None
    end = _parse_timestamp(end_date).date() if end_date else None

    if start and end and end < start:
        raise ValueError("End date cannot be earlier than start date.")

    selected = [
        record
        for record in records
        if (start is None or record["timestamp"].date() >= start)
        and (end is None or record["timestamp"].date() <= end)
    ]

    if not selected:
        raise ValueError("No usage records exist in the selected period.")

    return selected


def calculate_flat_rate(consumption, rate, fixed_fee=0):
    """
    Calculate an electricity bill using a flat rate tariff.

    Parameters:
        - consumption (float): Electricity consumed in kWh.
        - rate (float): Price per kWh.
        - fixed_fee (float): Fixed supply fee.

    Returns:
        - float: Total electricity bill.

    Raises:
        - TypeError: If consumption, rate, or fixed_fee is not a number.
        - ValueError: If consumption, rate, or fixed_fee is negative.
    """
    consumption = _number("consumption", consumption)
    rate = _number("rate", rate)
    fixed_fee = _number("fixed_fee", fixed_fee)

    energy_cost = consumption * rate
    total_bill = energy_cost + fixed_fee

    return round(total_bill, 2)


def calculate_flat_rate_breakdown(consumption, rate, fixed_fee=0):
    """Return usage, energy charge, fixed fee, and total for a flat tariff."""
    consumption = _number("consumption", consumption)
    rate = _number("rate", rate)
    fixed_fee = _number("fixed_fee", fixed_fee)

    energy = round(consumption * rate, 2)

    return {
        "usage_kwh": round(consumption, 2),
        "energy_fee": energy,
        "fixed_fee": round(fixed_fee, 2),
        "total": round(energy + fixed_fee, 2),
    }


def _in_period(value, start, end):
    """Return whether a time is inside a normal or midnight-crossing period."""
    if start == end:
        return False

    if start < end:
        return start <= value < end

    return value >= start or value < end


def _validate_tou_periods(peak_start, peak_end, offpeak_start, offpeak_end):
    """Reject overlapping peak and off-peak periods."""
    sample_minutes = [
        time(hour, minute)
        for hour in range(24)
        for minute in range(60)
    ]

    if any(
        _in_period(value, peak_start, peak_end)
        and _in_period(value, offpeak_start, offpeak_end)
        for value in sample_minutes
    ):
        raise ValueError("Peak and off-peak periods cannot overlap.")


def calculate_tou_tariff(
    records,
    peak_rate,
    offpeak_rate,
    shoulder_rate,
    fixed_fee=0,
    peak_start=time(18, 0),
    peak_end=time(22, 0),
    offpeak_start=time(22, 0),
    offpeak_end=time(7, 0),
):
    """Apply timestamp-based TOU rates and return a category breakdown."""
    if not records:
        raise ValueError("At least one usage record is required.")

    rates = {
        "peak": _number("peak_rate", peak_rate),
        "off_peak": _number("offpeak_rate", offpeak_rate),
        "shoulder": _number("shoulder_rate", shoulder_rate),
    }

    fixed_fee = _number("fixed_fee", fixed_fee)

    for name, value in (
        ("peak_start", peak_start),
        ("peak_end", peak_end),
        ("offpeak_start", offpeak_start),
        ("offpeak_end", offpeak_end),
    ):
        if not isinstance(value, time):
            raise TypeError(f"{name} must be a datetime.time value.")

    _validate_tou_periods(
        peak_start,
        peak_end,
        offpeak_start,
        offpeak_end,
    )

    usage = {
        "peak": 0.0,
        "off_peak": 0.0,
        "shoulder": 0.0,
    }

    for index, record in enumerate(records, start=1):
        if "timestamp" not in record or "kWh" not in record:
            raise ValueError(
                f"Usage record {index} must contain timestamp and kWh."
            )

        timestamp = _parse_timestamp(record["timestamp"])
        amount = _number(f"record {index} kWh", record["kWh"])
        current = timestamp.time()

        if _in_period(current, peak_start, peak_end):
            category = "peak"
        elif _in_period(current, offpeak_start, offpeak_end):
            category = "off_peak"
        else:
            category = "shoulder"

        usage[category] += amount

    breakdown = {}
    energy_fee = 0.0

    for category in ("peak", "off_peak", "shoulder"):
        cost = usage[category] * rates[category]
        energy_fee += cost

        breakdown[category] = {
            "usage_kwh": round(usage[category], 2),
            "rate": round(rates[category], 4),
            "cost": round(cost, 2),
        }

    breakdown.update({
        "energy_fee": round(energy_fee, 2),
        "fixed_fee": round(fixed_fee, 2),
        "total": round(energy_fee + fixed_fee, 2),
    })

    return breakdown


def calculate_tiered_tariff(
    consumption,
    tier1_limit,
    tier1_rate,
    tier2_limit,
    tier2_rate,
    tier3_rate,
    fixed_fee=0,
):
    """
    Calculate an electricity bill using a progressive three-tier tariff.

    - Tier 1: Usage from 0 kWh up to tier1_limit.
    - Tier 2: Usage above tier1_limit up to tier2_limit.
    - Tier 3: Usage above tier2_limit.

    Parameters:
        - consumption (float): Total electricity consumption in kWh.
        - tier1_limit (float): Upper limit of Tier 1.
        - tier1_rate (float): Price per kWh for Tier 1.
        - tier2_limit (float): Upper limit of Tier 2.
        - tier2_rate (float): Price per kWh for Tier 2.
        - tier3_rate (float): Price per kWh for Tier 3.
        - fixed_fee (float): Fixed supply fee.

    Returns:
        - float: Total electricity bill.

    Raises:
        - TypeError: If a numeric input is not a number.
        - ValueError: If input values are negative or tier limits are invalid.
    """
    return calculate_tiered_breakdown(
        consumption,
        tier1_limit,
        tier1_rate,
        tier2_limit,
        tier2_rate,
        tier3_rate,
        fixed_fee,
    )["total"]


def calculate_tiered_breakdown(
    consumption,
    tier1_limit,
    tier1_rate,
    tier2_limit,
    tier2_rate,
    tier3_rate,
    fixed_fee=0,
):
    """Return per-tier usage/cost plus energy, fixed, and total charges."""
    consumption = _number("consumption", consumption)
    tier1_limit = _number("tier1_limit", tier1_limit)
    tier2_limit = _number("tier2_limit", tier2_limit)

    rates = [
        _number("tier1_rate", tier1_rate),
        _number("tier2_rate", tier2_rate),
        _number("tier3_rate", tier3_rate),
    ]

    fixed_fee = _number("fixed_fee", fixed_fee)

    if tier2_limit <= tier1_limit:
        raise ValueError(
            "Tier 2 limit must be greater than Tier 1 limit."
        )

    quantities = [
        min(consumption, tier1_limit),
        max(min(consumption, tier2_limit) - tier1_limit, 0),
        max(consumption - tier2_limit, 0),
    ]

    costs = [
        quantities[index] * rates[index]
        for index in range(3)
    ]

    energy_fee = sum(costs)

    return {
        "tier1": {
            "usage_kwh": round(quantities[0], 2),
            "rate": rates[0],
            "cost": round(costs[0], 2),
        },
        "tier2": {
            "usage_kwh": round(quantities[1], 2),
            "rate": rates[1],
            "cost": round(costs[1], 2),
        },
        "tier3": {
            "usage_kwh": round(quantities[2], 2),
            "rate": rates[2],
            "cost": round(costs[2], 2),
        },
        "energy_fee": round(energy_fee, 2),
        "fixed_fee": round(fixed_fee, 2),
        "total": round(energy_fee + fixed_fee, 2),
    }


def compare_tariffs(bills):
    """Return sorted bill totals, cheapest tariff, and savings against each option."""
    if not isinstance(bills, dict) or len(bills) < 2:
        raise ValueError(
            "At least two tariff bills are required for comparison."
        )

    checked = {
        name: _number(f"{name} bill", value)
        for name, value in bills.items()
    }

    cheapest = min(checked, key=checked.get)
    cheapest_total = checked[cheapest]

    return {
        "cheapest": cheapest,
        "cheapest_total": round(cheapest_total, 2),
        "bills": dict(
            sorted(checked.items(), key=lambda item: item[1])
        ),
        "savings": {
            name: round(value - cheapest_total, 2)
            for name, value in checked.items()
        },
    }


def generate_saving_suggestions(records, tou_breakdown, comparison):
    """Generate transparent suggestions from calculated usage and bill results."""
    if not records:
        raise ValueError("At least one usage record is required.")

    if "cheapest" not in comparison or "savings" not in comparison:
        raise ValueError("A valid tariff comparison is required.")

    suggestions = [
        f"The calculated lowest-cost option is {comparison['cheapest']} at ${comparison['cheapest_total']:.2f} for the selected period."
    ]

    peak_usage = tou_breakdown["peak"]["usage_kwh"]
    total_usage = sum(record["kWh"] for record in records)

    if total_usage and peak_usage / total_usage >= 0.20:
        suggestions.append(
            "Peak usage is at least 20% of total usage. Where practical, move flexible appliance use outside the configured peak period."
        )
    else:
        suggestions.append(
            "Peak usage is below 20% of total usage. Continue limiting discretionary electricity use during the configured peak period."
        )

    suggestions.append(
        "These estimates use the entered rates and usage data only; confirm tariff terms and fees with the electricity provider before changing plans."
    )

    return suggestions